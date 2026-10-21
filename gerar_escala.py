# -*- coding: utf-8 -*-
"""Gerador de escala com ausências por bloco e redução automática de tarefas.

Tenta a escala completa; se nenhuma combinação respeitar as regras,
mantém Cozinha e Sala e copa nos dois blocos, Panos só no primeiro
e o máximo possível de tarefas opcionais.
Os arquivos do site e os horários de confirmação continuam compatíveis.
"""
import os
import json
import random
import itertools
import subprocess
import shutil
import sys
import unicodedata
from collections import Counter
from datetime import datetime, timedelta

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

PESSOAS = ["JUKEBOX", "LOTERYA", "KOMIXÃO", "NAMOITA", "NAZARÉ", "CAMILA", "BAQUETADA", "AMANDA", "BELA"]
TAREFAS_INICIO = ["Área frontal da casa", "Sala e copa", "Cozinha", "Geladeira", "Lavanderia", "Panos"]
TAREFAS_FIM = ["Cozinha", "Armários cozinha", "Área traseira da casa", "Garagem", "Sala e copa"]
TAREFAS_OBRIGATORIAS = (("Cozinha", "Sala e copa", "Panos"), ("Cozinha", "Sala e copa"))
TAREFAS_DIFICEIS = {"Cozinha", "Armários cozinha", "Geladeira"}
PROIBIDO_PANOS = {"NAZARÉ", "CAMILA"}
BLOCOS = ("dom_seg_ter", "qui_sex_sab")
ROTULOS_BLOCOS = ("DOM/SEG/TER", "QUI/SEX/SÁB")
SEM_TAREFA = {"", "FOLGA", "AUSENTE"}
PASTA_SCRIPT = os.path.dirname(os.path.abspath(__file__))
ARQUIVO_HISTORICO = os.path.join(PASTA_SCRIPT, "historico_limpeza.csv")
ARQUIVO_JSON = os.path.join(PASTA_SCRIPT, "escalas.json")
GIT_REMOTE = "origin"
GIT_BRANCH = "main"
FAZER_PUSH_AUTOMATICO = True
ULTIMA_GERACAO = {}

# Base histórica usada APENAS se faltar uma pessoa em 06/09 no CSV e no JSON.
# Registros já existentes no CSV nunca são substituídos por esta base.
ESCALA_FIXA_06_09 = {
 "JUKEBOX": {"DOM/SEG/TER": "Área frontal da casa", "QUI/SEX/SÁB": "Cozinha"},
 "LOTERYA": {"DOM/SEG/TER": "Sala e copa", "QUI/SEX/SÁB": "Armários cozinha"},
 "KOMIXÃO": {"DOM/SEG/TER": "Cozinha", "QUI/SEX/SÁB": "FOLGA"},
 "NAMOITA": {"DOM/SEG/TER": "Geladeira", "QUI/SEX/SÁB": "Área traseira da casa"},
 "NAZARÉ": {"DOM/SEG/TER": "Lavanderia", "QUI/SEX/SÁB": "FOLGA"},
 "CAMILA": {"DOM/SEG/TER": "FOLGA", "QUI/SEX/SÁB": "Garagem"},
 "BAQUETADA": {"DOM/SEG/TER": "Panos", "QUI/SEX/SÁB": "Sala e copa"},
 "AMANDA": {"DOM/SEG/TER": "FOLGA", "QUI/SEX/SÁB": "FOLGA"},
 "BELA": {"DOM/SEG/TER": "FOLGA", "QUI/SEX/SÁB": "FOLGA"},
}



def encontrar_git():
    git = shutil.which("git")
    if git:
        return git
    caminhos_possiveis = [
        r"C:\Program Files\Git\cmd\git.exe",
        r"C:\Program Files\Git\bin\git.exe",
        r"C:\Program Files (x86)\Git\cmd\git.exe",
        r"C:\Program Files (x86)\Git\bin\git.exe",
    ]
    for caminho in caminhos_possiveis:
        if os.path.exists(caminho):
            return caminho
    raise FileNotFoundError(
        "O Git não foi encontrado no computador."
    )


def caminho_xlsx(data_semana):
    return os.path.join(
        PASTA_SCRIPT,
        f"escala_limpeza_{data_semana.strftime('%d_%m_%Y')}.xlsx",
    )


def carregar_historico():
    colunas = [
        "semana",
        "pessoa",
        "dom_seg_ter",
        "qui_sex_sab",
    ]
    if not os.path.exists(ARQUIVO_HISTORICO):
        return pd.DataFrame(columns=colunas)
    historico = pd.read_csv(
        ARQUIVO_HISTORICO,
        encoding="utf-8-sig",
    )
    for coluna in colunas:
        if coluna not in historico.columns:
            historico[coluna] = ""
    return historico[colunas].fillna("")


def salvar_historico(historico):
    historico.to_csv(
        ARQUIVO_HISTORICO,
        index=False,
        encoding="utf-8-sig",
    )


def ordenar_historico(historico):
    if historico.empty:
        return historico
    temp = historico.copy()
    temp["_data"] = pd.to_datetime(
        temp["semana"],
        format="%d/%m/%Y",
        errors="coerce",
    )
    ordem = {
        pessoa: i
        for i, pessoa in enumerate(PESSOAS)
    }
    temp["_ordem_pessoa"] = temp["pessoa"].map(ordem)
    temp = temp.sort_values([
        "_data",
        "_ordem_pessoa",
    ])
    return (
        temp
        .drop(columns=["_data", "_ordem_pessoa"])
        .reset_index(drop=True)
    )


def historico_antes_da_data(historico, data_nova):
    if historico.empty:
        return historico.copy()
    datas = pd.to_datetime(
        historico["semana"],
        format="%d/%m/%Y",
        errors="coerce",
    )
    return historico[datas < data_nova].copy()


def garantir_semana_fixa(historico):
    """Completa a semana-base sem sobrescrever nenhuma linha já registrada."""
    data_fixa = "06/09/2026"
    existentes = set(historico.loc[historico['semana'] == data_fixa, 'pessoa'])
    faltantes = [p for p in PESSOAS if p not in existentes]
    if not faltantes:
        return ordenar_historico(historico)
    base_json = carregar_json_site().get('2026-09-06', {}).get('pessoas', {})
    registros = []
    for pessoa in faltantes:
        base = base_json.get(pessoa, {})
        registros.append({
            'semana': data_fixa, 'pessoa': pessoa,
            'dom_seg_ter': base.get('dom_seg_ter', ESCALA_FIXA_06_09[pessoa]['DOM/SEG/TER']),
            'qui_sex_sab': base.get('qui_sex_sab', ESCALA_FIXA_06_09[pessoa]['QUI/SEX/SÁB']),
        })
    return ordenar_historico(pd.concat([historico, pd.DataFrame(registros)], ignore_index=True))


def normalizar_nome(nome):
    return ''.join(c for c in unicodedata.normalize('NFD', nome.strip().upper())
                   if unicodedata.category(c) != 'Mn')


def validar_ausentes(nomes):
    mapa = {normalizar_nome(p): p for p in PESSOAS}
    resultado = set()
    for nome in nomes or []:
        chave = normalizar_nome(nome)
        if not chave:
            continue
        if chave not in mapa:
            raise ValueError(f"Nome inválido: {nome}. Use: {', '.join(PESSOAS)}")
        resultado.add(mapa[chave])
    return resultado


def contexto_historico(historico, data_semana=None):
    contexto = {}
    historico = ordenar_historico(historico)
    semana_anterior = ((data_semana - timedelta(days=7)).strftime('%d/%m/%Y')
                       if data_semana is not None else
                       (historico.iloc[-1]['semana'] if not historico.empty else None))
    for pessoa in PESSOAS:
        linhas = historico[historico['pessoa'] == pessoa].to_dict('records')
        semanas = [r for r in linhas if any(r[b] not in ('', 'AUSENTE') for b in BLOCOS)]
        eventos = [r[b] for r in linhas for b in BLOCOS if r[b] != '']
        trabalhos = [t for t in eventos if t not in SEM_TAREFA]
        ultima = semanas[-1] if semanas else None
        calendario = next((r for r in reversed(linhas) if r['semana'] == semana_anterior), None)
        contexto[pessoa] = {
            'ultimo_evento': eventos[-1] if eventos else None,
            'ultimo_trabalho': trabalhos[-1] if trabalhos else None,
            'contagens': Counter(trabalhos), 'folgas': eventos.count('FOLGA'),
            'total': len(trabalhos),
            'dificeis': sum(t in TAREFAS_DIFICEIS for t in trabalhos),
            'recente1': {ultima[b] for b in BLOCOS} if ultima else set(),
            'recente2': {r[b] for r in semanas[-2:] for b in BLOCOS},
            'folga_anterior': bool(calendario and any(calendario[b] == 'FOLGA' for b in BLOCOS)),
            # Só a semana imediatamente anterior determina a folga semanal.
            # Não transfere uma obrigação de antes de uma semana AUSENTE.
            'folga_obrigatoria': bool(calendario and all(calendario[b] not in SEM_TAREFA for b in BLOCOS)),
        }
    return contexto


def buscar_escala(presentes_inicio, presentes_fim, tarefas_inicio, tarefas_fim, contexto, obrigatorias, data_semana, nao_podem_panos=None):
    """Busca com retrocesso: só reduz tarefas depois de esgotar o modo completo."""
    bloqueados_panos = PROIBIDO_PANOS | validar_ausentes(nao_podem_panos)
    presentes = (set(presentes_inicio), set(presentes_fim))
    tarefas = (tuple(tarefas_inicio), tuple(tarefas_fim))
    obrigatorias = set(obrigatorias)
    obrigatorias_inicio = {p for p in presentes[0] if contexto[p]['ultimo_evento'] in TAREFAS_DIFICEIS}
    impossiveis_fim = set()

    def ultimo_trabalho(pessoa, anterior):
        tarefa = anterior.get(pessoa)
        return tarefa if tarefa not in SEM_TAREFA and tarefa is not None else contexto[pessoa]['ultimo_trabalho']

    def opcoes_folga(indice, anterior):
        disponiveis = presentes[indice]
        vagas = len(disponiveis) - len(tarefas[indice])
        if vagas < 0:
            return []
        if indice == 0:
            necessarias = obrigatorias_inicio
            minimo_semanais = max(0, len(obrigatorias) - max(0, len(presentes[1]) - len(tarefas[1])))
        else:
            necessarias = {p for p in disponiveis if anterior.get(p) in TAREFAS_DIFICEIS}
            necessarias |= {p for p in obrigatorias if anterior.get(p) != 'FOLGA'}
            minimo_semanais = 0
        if not necessarias <= disponiveis or len(necessarias) > vagas:
            return []
        livres = sorted(disponiveis - necessarias)
        opcoes = []
        for extra in itertools.combinations(livres, vagas - len(necessarias)):
            folgas = necessarias | set(extra)
            if indice == 0 and len(folgas & obrigatorias) < minimo_semanais:
                continue
            prioridade = (
                -len(folgas & obrigatorias),
                sum(anterior.get(p) == 'FOLGA' for p in folgas),
                sum(contexto[p]['folga_anterior'] for p in folgas),
                sum(contexto[p]['folgas'] for p in folgas), random.random(),
            )
            opcoes.append((prioridade, folgas))
        opcoes.sort(key=lambda item: item[0])
        return [folgas for _, folgas in opcoes]

    def distribuir(indice, anterior):
        assinatura = tuple('D' if anterior.get(p) in TAREFAS_DIFICEIS else
                           'F' if anterior.get(p) == 'FOLGA' else
                           'A' if anterior.get(p) == 'AUSENTE' else 'L' for p in PESSOAS)
        if indice == 1 and assinatura in impossiveis_fim:
            return None
        for folgas in opcoes_folga(indice, anterior):
            escala = {p: ('AUSENTE' if p not in presentes[indice] else 'FOLGA') for p in PESSOAS}
            trabalhando = presentes[indice] - folgas

            def candidatos(tarefa, livres):
                nomes = [p for p in PESSOAS if p in livres]
                if tarefa == 'Panos':
                    nomes = [p for p in nomes if p not in bloqueados_panos]
                if tarefa in TAREFAS_DIFICEIS:
                    nomes = [p for p in nomes if ultimo_trabalho(p, anterior) not in TAREFAS_DIFICEIS]
                if indice == 0 and data_semana == datetime(2026, 9, 13):
                    nomes = [p for p in nomes if (p == 'KOMIXÃO') == (tarefa == 'Panos')]
                def prioridade(p):
                    c = contexto[p]
                    prioridade_dificil = 0
                    if tarefa in TAREFAS_DIFICEIS:
                        prioridade_dificil = int(not (
                            anterior.get(p) == 'FOLGA' if indice == 1 else c['folga_anterior']
                        ))
                    return (prioridade_dificil, anterior.get(p) == tarefa,
                            tarefa in c['recente1'], tarefa in c['recente2'], c['contagens'][tarefa],
                            c['dificeis'] if tarefa in TAREFAS_DIFICEIS else 0,
                            c['total'], random.random())
                return sorted(nomes, key=prioridade)

            def atribuir(restantes, livres):
                if not restantes:
                    if indice == 0:
                        segundo = distribuir(1, escala)
                        return (escala.copy(), segundo) if segundo is not None else None
                    return escala.copy()
                possibilidades = [(t, candidatos(t, livres)) for t in restantes]
                tarefa, nomes = min(possibilidades, key=lambda item: (len(item[1]), item[0] not in TAREFAS_DIFICEIS))
                if not nomes:
                    return None
                outras = tuple(t for t in restantes if t != tarefa)
                for pessoa in nomes:
                    escala[pessoa] = tarefa
                    # Só há duas vagas de folga no fim? Não escolha mais difíceis
                    # do que essas vagas permitem, contando as folgas semanais.
                    if indice == 0:
                        folgas_fim = {p for p in presentes[1] if escala[p] in TAREFAS_DIFICEIS}
                        folgas_fim |= {p for p in obrigatorias if p not in folgas}
                        if len(folgas_fim) > len(presentes[1]) - len(tarefas[1]):
                            escala[pessoa] = 'FOLGA'
                            continue
                    resultado = atribuir(outras, livres - {pessoa})
                    if resultado is not None:
                        return resultado
                    escala[pessoa] = 'FOLGA'
                return None

            resultado = atribuir(tarefas[indice], trabalhando)
            if resultado is not None:
                return resultado
        if indice == 1:
            impossiveis_fim.add(assinatura)
        return None

    return distribuir(0, {})


def validar_escala(escalas, ausencias, tarefas_blocos, contexto, obrigatorias, nao_podem_panos=None):
    """Validação final independente antes de permitir gravação/publicação."""
    bloqueados_panos = PROIBIDO_PANOS | validar_ausentes(nao_podem_panos)
    for indice, (escala, ausentes, esperadas) in enumerate(zip(escalas, ausencias, tarefas_blocos)):
        if set(escala) != set(PESSOAS):
            raise ValueError('A escala não contém todas as pessoas.')
        reais = [t for t in escala.values() if t not in SEM_TAREFA]
        if Counter(reais) != Counter(esperadas):
            raise ValueError(f'Tarefas incompletas no bloco {ROTULOS_BLOCOS[indice]}.')
        for p in PESSOAS:
            tarefa = escala[p]
            if (tarefa == 'AUSENTE') != (p in ausentes):
                raise ValueError(f'Ausência incorreta para {p}.')
            if tarefa == 'Panos' and p in bloqueados_panos:
                raise ValueError(f'{p} não pode receber Panos.')
            anterior = contexto[p]['ultimo_trabalho']
            if indice == 1 and escalas[0][p] not in SEM_TAREFA:
                anterior = escalas[0][p]
            if tarefa in TAREFAS_DIFICEIS and anterior in TAREFAS_DIFICEIS:
                raise ValueError(f'Duas tarefas difíceis seguidas para {p}.')
            if indice == 0 and p not in ausentes and contexto[p]['ultimo_evento'] in TAREFAS_DIFICEIS and tarefa != 'FOLGA':
                raise ValueError(f'Falta folga após tarefa difícil de {p}.')
            if indice == 1 and escalas[0][p] in TAREFAS_DIFICEIS and p not in ausentes and tarefa != 'FOLGA':
                raise ValueError(f'Falta folga no segundo bloco para {p}.')
    for p in obrigatorias:
        if all(escala[p] != 'FOLGA' for escala in escalas):
            raise ValueError(f'Falta a folga semanal obrigatória de {p}.')


def gerar_semana(data_semana, ausentes_inicio, ausentes_fim=None, nao_podem_panos=None):
    global ULTIMA_GERACAO
    ULTIMA_GERACAO = {}
    if data_semana.weekday() != 6:
        raise ValueError('A data informada não é um domingo.')
    ausencias = (validar_ausentes(ausentes_inicio), validar_ausentes(ausentes_fim))
    nao_podem_panos = validar_ausentes(nao_podem_panos)
    historico = garantir_semana_fixa(carregar_historico())
    if data_semana == datetime(2026, 9, 6):
        if any(ausencias):
            raise ValueError('06/09/2026 é uma semana histórica fixa. Não altere as presenças dessa semana.')
        fixa = historico[historico['semana'] == '06/09/2026'].set_index('pessoa')
        ULTIMA_GERACAO = {'modo': 'histórico fixo'}
        return fixa['dom_seg_ter'].to_dict(), fixa['qui_sex_sab'].to_dict(), historico, []
    texto = data_semana.strftime('%d/%m/%Y')
    historico = historico[historico['semana'] != texto].copy()
    contexto = contexto_historico(historico_antes_da_data(historico, data_semana), data_semana)
    presentes = tuple([p for p in PESSOAS if p not in bloco] for bloco in ausencias)
    if not (set(presentes[0]) - PROIBIDO_PANOS - nao_podem_panos):
        raise ValueError('Nenhuma pessoa presente em DOM/SEG/TER pode fazer Panos nesta semana. '
                         'Confira os impedimentos de Panos e as ausências. Nenhuma escala foi salva.')
    for indice, nomes in enumerate(presentes):
        if len(nomes) < len(TAREFAS_OBRIGATORIAS[indice]):
            raise ValueError(f'{ROTULOS_BLOCOS[indice]} tem {len(nomes)} pessoa(s) presente(s). '
                             f'São necessárias pelo menos {len(TAREFAS_OBRIGATORIAS[indice])} para '
                             + ', '.join(TAREFAS_OBRIGATORIAS[indice]) + ', sem acumular tarefas.')
    if data_semana == datetime(2026, 9, 13) and 'KOMIXÃO' in ausencias[0]:
        raise ValueError('KOMIXÃO precisa executar Panos em DOM/SEG/TER de 13/09/2026.')
    # Uma ausência no bloco já dispensa a pessoa de trabalhar nele;
    # não exigimos uma folga extra no único bloco em que ela comparece.
    obrigatorias = [p for p in PESSOAS if contexto[p]['folga_obrigatoria']
                    and p in presentes[0] and p in presentes[1]]
    # Enumera todas as combinações de opcionais, mantendo as essenciais de cada bloco.
    # Ordena pela quantidade TOTAL: só descarta uma tarefa quando nenhuma
    # combinação maior consegue cumprir presenças, folgas e restrições.
    opcionais = [[t for t in base if t not in TAREFAS_OBRIGATORIAS[i]]
                 for i, base in enumerate((TAREFAS_INICIO, TAREFAS_FIM))]
    combinacoes = []
    for n0 in range(len(opcionais[0]), -1, -1):
        for n1 in range(len(opcionais[1]), -1, -1):
            for extra0 in itertools.combinations(opcionais[0], n0):
                for extra1 in itertools.combinations(opcionais[1], n1):
                    tarefas = tuple([t for t in base if t in TAREFAS_OBRIGATORIAS[i] or t in extras]
                                    for i, (base, extras) in enumerate(zip((TAREFAS_INICIO, TAREFAS_FIM), (extra0, extra1))))
                    combinacoes.append(tarefas)
    combinacoes.sort(key=lambda t: -(len(t[0]) + len(t[1])))
    for tarefas in combinacoes:
        modo = 'completo' if tarefas == (TAREFAS_INICIO, TAREFAS_FIM) else 'adaptado'
        resultado = buscar_escala(presentes[0], presentes[1], tarefas[0], tarefas[1], contexto, obrigatorias, data_semana, nao_podem_panos)
        if resultado is None:
            continue
        validar_escala(resultado, ausencias, tarefas, contexto, obrigatorias, nao_podem_panos)
        ULTIMA_GERACAO = {
            'modo': modo,
            'impedidos_panos': [p for p in PESSOAS if p in nao_podem_panos],
            'motivo_reducao': ('opcionais retiradas: ' + '; '.join(ROTULOS_BLOCOS[i] + ': ' + ', '.join(t for t in base if t not in tarefas[i]) for i, base in enumerate((TAREFAS_INICIO, TAREFAS_FIM)) if any(t not in tarefas[i] for t in base))) if modo == 'adaptado' else '',
        }
        return resultado[0], resultado[1], historico, obrigatorias
    raise ValueError('Nem as tarefas obrigatórias têm combinação válida com as presenças e folgas obrigatórias. '
                     'Confira os ausentes de cada bloco. Nenhuma escala foi salva; as regras não foram ignoradas.')



def salvar_semana_historico(
    historico,
    data_semana,
    escala1,
    escala2,
):
    data_texto = data_semana.strftime("%d/%m/%Y")
    historico = historico[
        historico["semana"] != data_texto
    ].copy()
    registros = []
    for pessoa in PESSOAS:
        registros.append({
            "semana": data_texto,
            "pessoa": pessoa,
            "dom_seg_ter": escala1[pessoa],
            "qui_sex_sab": escala2[pessoa],
        })
    historico = pd.concat(
        [historico, pd.DataFrame(registros)],
        ignore_index=True,
    )
    historico = ordenar_historico(historico)
    salvar_historico(historico)
    return historico


def carregar_json_site():
    if not os.path.exists(ARQUIVO_JSON):
        return {}
    try:
        with open(
            ARQUIVO_JSON,
            "r",
            encoding="utf-8",
        ) as arquivo:
            return json.load(arquivo)
    except json.JSONDecodeError as erro:
        raise ValueError("O arquivo escalas.json está inválido. Corrija-o antes de salvar a escala.") from erro


def atualizar_json_site(
    data_semana,
    escala1,
    escala2,
):
    escalas = carregar_json_site()
    chave = data_semana.strftime("%Y-%m-%d")
    escalas[chave] = {
        "atualizada_em": datetime.now().strftime(
            "%d/%m/%Y %H:%M"
        ),
        "pessoas": {
            pessoa: {
                "dom_seg_ter": escala1[pessoa],
                "qui_sex_sab": escala2[pessoa],
            }
            for pessoa in PESSOAS
        },
    }
    escalas = dict(
        sorted(escalas.items())
    )
    with open(
        ARQUIVO_JSON,
        "w",
        encoding="utf-8",
    ) as arquivo:
        json.dump(
            escalas,
            arquivo,
            ensure_ascii=False,
            indent=2,
        )


def salvar_xlsx(
    data_semana,
    escala1,
    escala2,
):
    arquivo = caminho_xlsx(data_semana)
    wb = Workbook()
    ws = wb.active
    ws.title = data_semana.strftime("%d-%m-%Y")
    AZUL_ESCURO = "0B5A93"
    AZUL_QUARTA = "9DBBE8"
    VERMELHO = "FF0000"
    CINZA = "BFBFBF"
    BRANCO = "FFFFFF"
    PRETO = "000000"
    borda_fina = Side(
        style="thin",
        color=PRETO,
    )
    borda = Border(
        left=borda_fina,
        right=borda_fina,
        top=borda_fina,
        bottom=borda_fina,
    )
    alinhamento = Alignment(
        horizontal="center",
        vertical="center",
    )
    fonte_branca = Font(
        name="Arial",
        size=11,
        color=BRANCO,
    )
    fonte_branca_negrito = Font(
        name="Arial",
        size=11,
        bold=True,
        color=BRANCO,
    )
    fonte_normal = Font(
        name="Arial",
        size=11,
        color=PRETO,
    )
    ws.column_dimensions["A"].width = 17
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 32
    ws.merge_cells("B1:D1")
    titulo = ws["B1"]
    titulo.value = (
        "TABELA DE TAREFAS - "
        + data_semana.strftime("%d/%m/%Y")
    )
    titulo.font = fonte_branca_negrito
    titulo.alignment = alinhamento
    for coluna in range(2, 5):
        celula = ws.cell(row=1, column=coluna)
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor=AZUL_ESCURO,
        )
        celula.border = borda
    cabecalhos = [
        "",
        "DOMINGO/SEG/TERÇA",
        "QUARTA",
        "QUINTA/SEXTA/SÁB",
    ]
    for coluna, texto in enumerate(
        cabecalhos,
        start=1,
    ):
        celula = ws.cell(
            row=2,
            column=coluna,
        )
        celula.value = texto
        celula.alignment = alinhamento
        celula.border = borda
        if coluna > 1:
            celula.fill = PatternFill(
                fill_type="solid",
                fgColor=AZUL_ESCURO,
            )
            celula.font = fonte_branca_negrito
    linha = 3
    for pessoa in PESSOAS:
        tarefa1 = escala1[pessoa]
        tarefa2 = escala2[pessoa]
        nome = ws.cell(
            row=linha,
            column=1,
        )
        nome.value = pessoa
        nome.fill = PatternFill(
            fill_type="solid",
            fgColor=AZUL_ESCURO,
        )
        nome.font = fonte_branca
        nome.alignment = alinhamento
        nome.border = borda
        c1 = ws.cell(row=linha, column=2)
        if tarefa1 == "FOLGA":
            c1.value = ""
            c1.fill = PatternFill(
                fill_type="solid",
                fgColor=VERMELHO,
            )
        elif tarefa1 == "AUSENTE":
            c1.value = "AUSENTE"
            c1.fill = PatternFill(
                fill_type="solid",
                fgColor=CINZA,
            )
        else:
            c1.value = tarefa1
            c1.fill = PatternFill(
                fill_type="solid",
                fgColor=BRANCO,
            )
        c1.font = fonte_normal
        c1.alignment = alinhamento
        c1.border = borda
        c2 = ws.cell(row=linha, column=3)
        c2.value = ""
        c2.fill = PatternFill(
            fill_type="solid",
            fgColor=AZUL_QUARTA,
        )
        c2.border = borda
        c3 = ws.cell(row=linha, column=4)
        if tarefa2 == "FOLGA":
            c3.value = ""
            c3.fill = PatternFill(
                fill_type="solid",
                fgColor=VERMELHO,
            )
        elif tarefa2 == "AUSENTE":
            c3.value = "AUSENTE"
            c3.fill = PatternFill(
                fill_type="solid",
                fgColor=CINZA,
            )
        else:
            c3.value = tarefa2
            c3.fill = PatternFill(
                fill_type="solid",
                fgColor=BRANCO,
            )
        c3.font = fonte_normal
        c3.alignment = alinhamento
        c3.border = borda
        ws.row_dimensions[linha].height = 22
        linha += 1
    ws.sheet_view.showGridLines = False
    wb.save(arquivo)
    return arquivo


def executar_git(*args):
    resultado = subprocess.run(
        [encontrar_git(), *args],
        cwd=PASTA_SCRIPT,
        text=True,
        capture_output=True
    )
    if resultado.returncode != 0:
        raise RuntimeError(
            resultado.stderr.strip()
            or resultado.stdout.strip()
            or "Erro ao executar Git."
        )
    return resultado.stdout.strip()


def sincronizar_github():
    # Atualiza a pasta local antes de gerar uma nova escala.
    # --autostash permite o rebase mesmo se houver alterações locais
    # em arquivos já rastreados pelo Git.
    executar_git(
        "rev-parse",
        "--is-inside-work-tree",
    )
    print("Sincronizando com o GitHub...")
    executar_git(
        "pull",
        "--rebase",
        "--autostash",
        GIT_REMOTE,
        GIT_BRANCH,
    )
    print("Git pull --rebase concluído.")


def publicar_github(data_semana):
    # Confere se estamos dentro de um repositório
    executar_git("rev-parse", "--is-inside-work-tree")
    # Adiciona somente os arquivos relevantes do projeto
    arquivos = [
        "historico_limpeza.csv",
        "escalas.json",
        "index.html",
        "style.css",
        "script.js",
        os.path.basename(__file__),
    ]
    xlsx_relativo = os.path.basename(
        caminho_xlsx(data_semana)
    )
    arquivos.append(xlsx_relativo)
    executar_git("add", *arquivos)
    # Verifica se existe alteração staged
    diff = subprocess.run(
        [
            encontrar_git(),
            "diff",
            "--cached",
            "--quiet"
        ],
        cwd=PASTA_SCRIPT,
        capture_output=True,
        text=True,
    )
    if diff.returncode not in (0, 1):
        raise RuntimeError(diff.stderr.strip() or "Erro ao verificar alterações do Git.")
    if diff.returncode == 0:
        print("Não há alterações novas para commit.")
        return
    mensagem = (
        "Atualiza escala "
        + data_semana.strftime("%d/%m/%Y")
    )
    executar_git(
        "commit",
        "-m",
        mensagem,
    )
    if FAZER_PUSH_AUTOMATICO:
        executar_git(
            "push",
            GIT_REMOTE,
            GIT_BRANCH,
        )
        print("Git push concluído com sucesso.")


def mostrar_escala(
    data_semana,
    escala1,
    escala2,
    folgas_obrigatorias,
):
    print()
    print("=" * 100)
    print(
        "TABELA DE TAREFAS - "
        + data_semana.strftime("%d/%m/%Y")
    )
    print("=" * 100)
    df = pd.DataFrame(
        {
            "DOM/SEG/TER": [
                escala1[pessoa]
                for pessoa in PESSOAS
            ],
            "QUI/SEX/SÁB": [
                escala2[pessoa]
                for pessoa in PESSOAS
            ],
        },
        index=PESSOAS,
    )
    print("Modo: " + ULTIMA_GERACAO.get("modo", "completo"))
    if ULTIMA_GERACAO.get("motivo_reducao"):
        print("Escala adaptada: " + ULTIMA_GERACAO["motivo_reducao"])
    impedidos = ULTIMA_GERACAO.get('impedidos_panos', [])
    if impedidos:
        print('Não podem fazer Panos nesta semana: ' + ', '.join(impedidos))
    print(df.to_string())
    if folgas_obrigatorias:
        print()
        print(
            "Folga obrigatória nesta semana: "
            + ", ".join(folgas_obrigatorias)
        )


def ler_ausentes(bloco):
    print(f'\nQuem estará AUSENTE em {bloco}?')
    print('Separe os nomes por vírgula; pressione Enter se todas estiverem presentes.')
    texto = input(f'Ausentes em {bloco}: ').strip()
    nomes = validar_ausentes(texto.split(','))
    return [p for p in PESSOAS if p in nomes]


def ler_impedidos_panos():
    print('\nQuem não pode fazer Panos nesta semana?')
    print('Separe os nomes por vírgula. Enter mantém somente as restrições permanentes.')
    print('Essas pessoas continuam disponíveis para outras tarefas no bloco.')
    print('NAZARÉ e CAMILA continuam proibidas de fazer Panos em qualquer semana.')
    nomes = validar_ausentes(input('Não podem fazer Panos: ').strip().split(','))
    return [p for p in PESSOAS if p in nomes]


def main():
    print('\nGERADOR DE ESCALA DE LIMPEZA')
    print('=' * 50)
    data_semana = datetime.strptime(input('Digite o domingo da nova escala (DD/MM/AAAA): ').strip(), '%d/%m/%Y')
    if data_semana.weekday() != 6:
        raise ValueError('A data informada não é um domingo.')
    ausentes_inicio = ler_ausentes('DOM/SEG/TER')
    ausentes_fim = ler_ausentes('QUI/SEX/SÁB')
    nao_podem_panos = ler_impedidos_panos()
    if FAZER_PUSH_AUTOMATICO:
        print()
        sincronizar_github()
    escala1, escala2, historico, folgas_obrigatorias = gerar_semana(data_semana, ausentes_inicio, ausentes_fim, nao_podem_panos)
    mostrar_escala(data_semana, escala1, escala2, folgas_obrigatorias)
    if input('\nDeseja salvar e publicar essa escala? (s/n): ').strip().lower() != 's':
        print('Escala não salva.')
        return
    # Valida a integridade do JSON antes de alterar qualquer arquivo de dados.
    carregar_json_site()
    salvar_semana_historico(historico, data_semana, escala1, escala2)
    arquivo_xlsx = salvar_xlsx(data_semana, escala1, escala2)
    atualizar_json_site(data_semana, escala1, escala2)
    print('\nArquivos locais atualizados:')
    print(f'- CSV:  {ARQUIVO_HISTORICO}')
    print(f'- JSON: {ARQUIVO_JSON}')
    print(f'- XLSX: {arquivo_xlsx}')
    if FAZER_PUSH_AUTOMATICO:
        print('\nPublicando no GitHub...')
        publicar_github(data_semana)
    print('\nConcluído.')


if __name__ == '__main__':
    try:
        main()
    except (ValueError, RuntimeError, FileNotFoundError, subprocess.SubprocessError) as erro:
        print('\nERRO:')
        print(erro)
        sys.exit(1)


